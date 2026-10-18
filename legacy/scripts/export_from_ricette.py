#!/usr/bin/env python3
"""Export the recipes of ../Ricette.xlsx (Google-Sheets workbook) to seed/recipes/*.json + seed/media/*.jpg
in the calderone-recipe/1 format.  Ids are uuid5(title) so re-exporting is stable and re-importing idempotent.

Usage: conda run -n CLAUDE python scripts/export_from_ricette.py [--xlsx ../Ricette.xlsx] [--parsed ../parsed] [--out seed]
"""
import argparse
import io
import json
import re
import uuid
from pathlib import Path

import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
NS = uuid.UUID("7d1f3c1e-5b0a-4c55-9b8e-ca1de7a00001")
# cooking times of the 4 original recipes (read from their app cover screenshots; the 5 newer ones come from parsed/)
KNOWN_TIMES = {"FUSILLI AL PESTO DI PISTACCHI E MENTA": 20, "ZUPPA DI ZUCCA CON QUINOA E SPINACINO": 40,
               "ZUPPA DI MISO": 30, "PASTA AL PESTO DI RADICCHIO E NOCI": 20}


def sentence_case(s: str) -> str:
    return s[:1].upper() + s[1:].lower()


def parse_sheet(ws):
    title = ws["A1"].value.strip()
    m = re.search(r"(\d+)\s*pax", str(ws["B21"].value or ""))
    servings = int(m.group(1)) if m else 4
    ingredients, r = [], 22
    while ws.cell(r, 1).value not in (None, ""):
        name, qty, unit = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value or ""
        ingredients.append(dict(name=str(name).strip(), qty=float(qty) if isinstance(qty, (int, float)) else None,
                                unit=str(unit).strip() or "q.b.", optional=False, note=""))
        r += 1
    steps, notes, pending, mode = [], "", None, None
    for row in range(r, ws.max_row + 1):
        v = ws.cell(row, 1).value
        if v is None or not str(v).strip():
            continue
        v = " ".join(str(v).split())
        if v == "Preparazione":
            mode = "steps"; continue
        if v == "Note":
            mode = "notes"; continue
        if mode == "notes":
            notes = (notes + " " + v).strip(); continue
        mt = re.match(r"^\d+/\d+\s+(.*)$", v)
        if mt:
            pending = dict(title=mt.group(1), text=""); steps.append(pending)
        elif pending is not None:
            pending["text"] = (pending["text"] + " " + v).strip()
    # "Cipolla e pepe: facoltativi." -> optional flags
    mo = re.match(r"^(.*?):\s*facoltativ\w*\.?\s*$", notes, re.I)
    if mo:
        names = {n.strip().lower() for n in re.split(r",| e ", mo.group(1))}
        for i in ingredients:
            if i["name"].lower() in names:
                i["optional"] = True
        notes = ""
    return title, servings, ingredients, steps, notes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(ROOT.parent / "Ricette.xlsx")); ap.add_argument("--out", default=str(ROOT / "seed"))
    ap.add_argument("--parsed", default=str(ROOT.parent / "parsed"), help="transcribed JSON (times, source screenshots)")
    a = ap.parse_args()
    out = Path(a.out); (out / "recipes").mkdir(parents=True, exist_ok=True); (out / "media").mkdir(exist_ok=True)
    parsed = {}
    for f in Path(a.parsed).glob("*.json"):
        d = json.loads(f.read_text(encoding="utf-8")); parsed[d["title"].upper()] = d
    wb = openpyxl.load_workbook(a.xlsx)
    cats = {wb["ARCHIVIO RICETTE"].cell(r, 1).value: wb["ARCHIVIO RICETTE"].cell(r, 3).value for r in range(2, 200)}
    first = wb.sheetnames.index("ARCHIVIO INGREDIENTI") + 1
    for ws in wb.worksheets[first:]:
        title, servings, ingredients, steps, notes = parse_sheet(ws)
        rid = uuid.uuid5(NS, title).hex
        p = parsed.get(title, {})
        photo = None
        if ws._images:
            photo = f"{rid}.jpg"
            Image.open(io.BytesIO(ws._images[0]._data())).convert("RGB").save(out / "media" / photo, "JPEG", quality=88)
        src = ["Ricette.xlsx"] + ([", ".join(Path(s).name for s in p["source_images"])] if p.get("source_images") else [])
        rec = {"format": "calderone-recipe/1", "id": rid, "title": sentence_case(title), "servings": servings,
               "time_minutes": p.get("time_minutes") or KNOWN_TIMES.get(title), "category": cats.get(title) or "Piatto unico",
               "tags": [], "ingredients": ingredients, "steps": steps, "notes": notes,
               "source": " · ".join(src), "photo": photo, "parents": []}
        (out / "recipes" / f"{rid}.json").write_text(json.dumps(rec, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"{rec['title'][:45]:45} {len(ingredients):2} ingr. {len(steps)} passi  {rec['time_minutes']} min  foto={bool(photo)}")


if __name__ == "__main__":
    main()
